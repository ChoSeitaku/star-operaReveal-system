"""附件解析：提取文本预览 → 写 sidecar 元数据 → 回写 attachments.parse_status。

解析在后台任务中执行，不阻塞上传响应。中文 CSV 常见 GBK/BOM 编码，
因此解码采用 utf-8-sig → utf-8 → gb18030 的三级回退。
"""
from __future__ import annotations

import asyncio
import json
from pathlib import Path
from typing import Any

from backend import config, database
from backend.utils import truncate

# 纯文本类扩展名
TEXT_EXTENSIONS = {"csv", "txt", "md", "json", "log", "tsv", "sql", "yaml", "yml"}
# 表格类扩展名（需 openpyxl）
SHEET_EXTENSIONS = {"xlsx", "xlsm"}
# 明确不支持（旧版二进制格式）
UNSUPPORTED_EXTENSIONS = {"xls"}

XLSX_MAX_SHEETS = 3
XLSX_MAX_ROWS = 200


class ParseError(Exception):
    """附件解析失败。"""


def decode_text(raw: bytes) -> tuple[str, str]:
    """按 utf-8-sig → utf-8 → gb18030 顺序解码，返回 (文本, 实际编码)。"""
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8(replace)"


def sidecar_path(file_path: str | Path) -> Path:
    return Path(f"{file_path}.meta.json")


def _extract_text(path: Path) -> dict[str, Any]:
    raw = path.read_bytes()
    text, encoding = decode_text(raw)
    preview_chars = config.get_int("upload.preview_chars", 5000)
    lines = text.count("\n") + 1
    return {
        "encoding": encoding,
        "char_count": len(text),
        "line_count": lines,
        "preview": truncate(text, preview_chars, suffix="\n…（预览已截断）"),
    }


def _extract_sheet(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ParseError("缺少 openpyxl 依赖，无法解析表格文件") from exc

    try:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError(f"表格文件无法打开: {exc}") from exc

    blocks: list[str] = []
    total_rows = 0
    try:
        for sheet_name in workbook.sheetnames[:XLSX_MAX_SHEETS]:
            sheet = workbook[sheet_name]
            blocks.append(f"### 工作表: {sheet_name}")
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= XLSX_MAX_ROWS:
                    blocks.append("…（本表仅预览前 %d 行）" % XLSX_MAX_ROWS)
                    break
                total_rows += 1
                cells = ["" if cell is None else str(cell) for cell in row]
                blocks.append("\t".join(cells).rstrip())
    finally:
        workbook.close()

    preview_chars = config.get_int("upload.preview_chars", 5000)
    preview = truncate("\n".join(blocks), preview_chars, suffix="\n…（预览已截断）")
    return {"encoding": "xlsx", "sheet_count": len(workbook.sheetnames), "row_count": total_rows, "preview": preview}


def extract(path: Path, file_type: str) -> dict[str, Any]:
    """同步解析入口（在线程中调用）。"""
    ext = (file_type or "").lower()
    if ext in UNSUPPORTED_EXTENSIONS:
        raise ParseError("不支持旧版 .xls 二进制格式，请另存为 .xlsx 后重新上传")
    if ext in SHEET_EXTENSIONS:
        return _extract_sheet(path)
    if ext in TEXT_EXTENSIONS:
        return _extract_text(path)
    raise ParseError(f"不支持的文件类型: .{ext}（支持 {', '.join(sorted(TEXT_EXTENSIONS | SHEET_EXTENSIONS))}）")


def write_sidecar(file_path: str | Path, meta: dict[str, Any]) -> None:
    try:
        sidecar_path(file_path).write_text(
            json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    except OSError:
        pass  # 元数据写入失败不影响主流程


async def parse_attachment(attachment_id: int) -> None:
    """后台解析任务：更新 parse_status 并落地 sidecar。"""
    row = await database.fetch_one("SELECT * FROM attachments WHERE id = ?", (attachment_id,))
    if row is None:
        return

    file_path = Path(row["file_path"])
    await database.execute("UPDATE attachments SET parse_status = 'parsing' WHERE id = ?", (attachment_id,))

    meta: dict[str, Any] = {"attachment_id": attachment_id, "file_name": row["file_name"]}
    try:
        if not file_path.exists():
            raise ParseError("文件不存在或已被删除")
        result = await asyncio.to_thread(extract, file_path, row["file_type"])
        meta.update(result)
        meta["parse_status"] = "success"
        meta["error"] = None
    except ParseError as exc:
        meta["parse_status"] = "failed"
        meta["error"] = str(exc)
    except Exception as exc:  # 兜底：任何异常都落到 failed，避免后台任务静默失败
        meta["parse_status"] = "failed"
        meta["error"] = f"解析异常: {exc}"

    await asyncio.to_thread(write_sidecar, file_path, meta)
    await database.execute(
        "UPDATE attachments SET parse_status = ? WHERE id = ?", (meta["parse_status"], attachment_id)
    )


def load_preview(file_path: str | Path, max_chars: int) -> tuple[str | None, str]:
    """读取 sidecar 中的预览文本，返回 (预览, 解析状态)。"""
    path = sidecar_path(file_path)
    if not path.exists():
        return None, "unknown"
    try:
        meta = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None, "unknown"
    status = str(meta.get("parse_status") or "unknown")
    preview = meta.get("preview")
    if not isinstance(preview, str):
        return None, status
    return truncate(preview, max_chars), status


def remove_sidecar(file_path: str | Path) -> None:
    try:
        sidecar_path(file_path).unlink(missing_ok=True)
    except OSError:
        pass
